# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import copy
import segno

# 在线试玩地址（扫码进入；如重新部署请同步更新此 URL）
PLAY_URL = "https://annabellexyq.github.io/youyi-tianxia-3d/"
QR_PATH  = "/Users/xuyueqing/CodeBuddy/20260903210153/youyi-tianxia-3d/qrcode_play.png"

# palette
INK   = RGBColor(0x2b,0x27,0x23)
INK2  = RGBColor(0x4a,0x44,0x3d)
PAPER = RGBColor(0xec,0xe7,0xda)
PAPER2= RGBColor(0xe3,0xdd,0xd0)
GOLD  = RGBColor(0x9c,0x7b,0x3f)
RED   = RGBColor(0xa3,0x3a,0x32)
LINE  = RGBColor(0xcd,0xc6,0xb6)
WHITE = RGBColor(0xf5,0xf1,0xe7)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

def bg(slide, color=PAPER):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color

def credit(slide):
    tb = box(slide, 10.6, 0.22, 2.4, 0.4)
    set_text(tb.text_frame, "制作人：AnnabelleXu", 11, RGBColor(0x9a,0x9a,0x9a), True, PP_ALIGN.RIGHT, 1, "STSong")

def box(slide, l,t,w,h):
    return slide.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))

def set_text(tf, text, size, color=INK, bold=False, align=PP_ALIGN.LEFT, spacing=1.1, font="STSong"):
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold
    r.font.color.rgb = color; r.font.name = font
    p.line_spacing = spacing
    return p

def add_para(tf, text, size, color=INK, bold=False, align=PP_ALIGN.LEFT, spacing=1.15, bullet=False, level=0, font="STSong"):
    segs = text.split('\n')
    for j, seg in enumerate(segs):
        p = tf.paragraphs[0] if (j == 0 and len(tf.paragraphs) == 1 and not tf.paragraphs[0].runs) else tf.add_paragraph()
        p.alignment = align; p.level = level
        r = p.add_run(); r.text = (("• " if (bullet and j == 0) else "") + seg)
        r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = color; r.font.name = font
        p.line_spacing = spacing
    return p

def tag(slide, text):
    set_text(box(slide,0.9,0.55,8,0.5).text_frame, text, 13, GOLD, True, spacing=1)
    slide.shapes[-1].text_frame.paragraphs[0].runs[0].font.name = "Arial"

def heading(slide, text):
    tb = box(slide,0.9,1.15,11,0.9)
    p = set_text(tb.text_frame, text, 30, INK, True, spacing=1)
    # red bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9),Inches(2.05),Inches(0.16),Inches(0.55))
    bar.fill.solid(); bar.fill.fore_color.rgb = RED; bar.line.fill.background()
    tb.left = Inches(1.25)

def card(slide, l,t,w,h, k, v):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l),Inches(t),Inches(w),Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = WHITE
    shp.line.color.rgb = LINE; shp.line.width = Pt(1)
    tf = shp.text_frame; tf.word_wrap = True
    tf.margin_left=Inches(0.18); tf.margin_right=Inches(0.15); tf.margin_top=Inches(0.14); tf.margin_bottom=Inches(0.12)
    set_text(tf, k, 16, RED, True, spacing=1.05)
    add_para(tf, v, 12, INK2, spacing=1.25, font="STSong")
    # left gold strip
    strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l),Inches(t),Inches(0.06),Inches(h))
    strip.fill.solid(); strip.fill.fore_color.rgb = GOLD; strip.line.fill.background()

# ---------------- Cover ----------------
s = prs.slides.add_slide(BLANK); bg(s, PAPER2)
set_text(box(s,1.0,2.0,11,1.6).text_frame, "游 医 天 下", 64, INK, True, PP_ALIGN.CENTER, 1)
set_text(box(s,1.0,3.5,11,0.8).text_frame, "白描水墨 · 五脏五色 · 中医诊疗叙事", 22, INK2, False, PP_ALIGN.CENTER, 1)
set_text(box(s,1.0,4.4,11,0.6).text_frame, "九境地图 · 类手工白描场景 · 望色归脏 · 五脏显影", 15, GOLD, True, PP_ALIGN.CENTER, 1)
# frame
fr = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45),Inches(0.45),Inches(12.43),Inches(6.6))
fr.fill.background(); fr.line.color.rgb = LINE; fr.line.width = Pt(1.5)

# ---------------- 01 产品定位 ----------------
s = prs.slides.add_slide(BLANK); bg(s); credit(s)
tag(s,"01 / POSITIONING")
heading(s,"产品定位")
items = [
 ("品类","以中医「五脏五色」学说为核心的国风叙事解谜游戏，将传统诊疗化为可玩机制。"),
 ("美术","白描水墨 + 低多边形立体，黑白山水意境与现代程序化美术相融合。"),
 ("形态","上帝视角九境地图探索，每境一座 Three.js 实时渲染的类手工白描场景，循河而上次第解锁。"),
 ("受众","国风文化爱好者、中医科普人群、叙事与解谜游戏玩家；兼具娱乐与文化传播价值。"),
 ("差异化","把「处方」变成可操作的五行调色系统，以望色归脏与五脏显影承载叙事，而非传统战斗驱动。"),
 ("价值主张","在游戏里读懂五脏五色，用「拟方」改变世界，以一介游医之眼遍历九境时疫。"),
]
pos=[(0.9,2.5),(6.9,2.5),(0.9,4.0),(6.9,4.0),(0.9,5.5),(6.9,5.5)]
for (k,v),(l,t) in zip(items,pos):
    card(s,l,t,5.5,1.25,k,v)

# ---------------- 02 核心技术应用 ----------------
s = prs.slides.add_slide(BLANK); bg(s); credit(s)
tag(s,"02 / TECHNOLOGY")
heading(s,"核心技术应用")
items = [
 ("Three.js 实时渲染","程序化占位 + glTF 加载，类手工白描场景的雾色、灯光随组方实时变化，构建可交互世界。"),
 ("glTF 资产加载","支持外部模型导入与程序化占位，人物与场景可平滑替换升级。"),
 ("五脏五色组方系统","数据驱动的状态映射：五色剂量实时作用于场景雾色、灯光与角色腰带色。"),
 ("望色归脏推演","拍照上传 → 五色归脏 → 拟方推演，并结合八字四柱参定先天体质与取向。"),
 ("手工级五脏显影","角色身上实时显影「五色入五脏」，把中医理论具象化为可视反馈。"),
 ("轻量前端架构","importmap + ES Modules，零构建即可在浏览器运行，便于演示与分发。"),
]
pos=[(0.9,2.5),(6.9,2.5),(0.9,4.0),(6.9,4.0),(0.9,5.5),(6.9,5.5)]
for (k,v),(l,t) in zip(items,pos):
    card(s,l,t,5.5,1.25,k,v)

# ---------------- 03 游戏创意点 ----------------
s = prs.slides.add_slide(BLANK); bg(s); credit(s)
tag(s,"03 / CREATIVITY")
heading(s,"游戏创意点说明")
items = [
 ("拟方即调色","将中医处方抽象为「肝青/心赤/脾黄/肺白/肾黑」五色剂量，用调色盘代替开方。"),
 ("五脏显影","角色身上实时显影「五色入五脏」的疗效反馈，具象化中医理论，可开合查看。"),
 ("望色归脏","上传一张照片，依「五色入五脏」推演归脏并拟方，把望诊变成可玩机制。"),
 ("环境即反馈","场景雾色、灯光、腰带色随组方实时晕染，让「药效」可见可感。"),
 ("九境叙事","上帝视角地图循河解锁，九境各有偏性，组方施治串起整段行医治验。"),
 ("白描 × 立体","类手工白描线稿质感叠加低多边形立体空间，形成独特的中式极简视觉语言。"),
]
pos=[(0.9,2.5),(6.9,2.5),(0.9,4.0),(6.9,4.0),(0.9,5.5),(6.9,5.5)]
for (k,v),(l,t) in zip(items,pos):
    card(s,l,t,5.5,1.25,k,v)

# ---------------- 04 核心玩法 ----------------
s = prs.slides.add_slide(BLANK); bg(s); credit(s)
tag(s,"04 / GAMEPLAY")
heading(s,"核心玩法说明")
steps = [
 ("壹","探图解锁","上帝视角九境地图，\n循河而上逐境解锁。"),
 ("贰","入境观察","进入类手工白描场景，\n观察环境与人物气色，\n判断该境偏性。"),
 ("叁","拟方调色","调五脏五色剂量，\n实时晕染雾色、灯光\n与角色腰带色。"),
 ("肆","五脏显影","查看角色身上\n五色入五脏的疗效反馈。"),
 ("伍","望色归脏","拍照上传，五色归脏\n并拟方推演，\n结合八字参定先天体质。"),
]
x=0.65; w=2.45; gap=0.13; top=2.7
for i,(n,t,d) in enumerate(steps):
    l = x + i*(w+gap)
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l),Inches(top),Inches(w),Inches(2.4))
    shp.fill.solid(); shp.fill.fore_color.rgb = WHITE; shp.line.color.rgb = GOLD; shp.line.width = Pt(1.25)
    tf=shp.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    tf.margin_top=Inches(0.18); tf.margin_bottom=Inches(0.18)
    set_text(tf, n, 18, GOLD, True, PP_ALIGN.CENTER, 1.1)
    add_para(tf, t, 17, RED, True, PP_ALIGN.CENTER, 1.2)
    add_para(tf, d, 10.5, INK2, False, PP_ALIGN.CENTER, 1.25)
    if i<len(steps)-1:
        ar = s.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(l+w-0.02),Inches(top+0.95),Inches(0.22),Inches(0.5))
        ar.fill.solid(); ar.fill.fore_color.rgb = GOLD; ar.line.fill.background()
set_text(box(s,0.9,5.5,11.5,1.3).text_frame,
         "玩家以「游医」身份遍历九境，用一套可视化的五脏五色组方体系应对不同疫病：\n每一次调色都是一次处方，每一处显影都是一次疗效，世界因你的药方而改变。",
         15, INK2, False, PP_ALIGN.LEFT, 1.3)

# ---------------- 05 GNPC 智能体对话 ----------------
s = prs.slides.add_slide(BLANK); bg(s); credit(s)
tag(s,"05 / NPC · GNPC")
heading(s,"会聊天的药灵 · 智能体对话")
items = [
 ("GNPC 智能体平台","接入腾讯云 GNPC 智能体，把五脏五色、性味归经等中医知识注入可对话 NPC。"),
 ("外泌体晨光","游戏内常驻 NPC，玩家随时点开与它对话，问医理、问剧情、问组方思路。"),
 ("望色归脏推演","上传照片后，它依「五色入五脏」给出归脏与拟方推演，而非空泛应答。"),
 ("对话驱动叙事","NPC 依剧情语境应答，让「游医行脚」的叙事随对话自然推进。"),
 ("同源安全调用","NPC 调用经本地代理签名转发，密钥不落前端，规避 CORS 与泄露。"),
 ("可玩化科普","把枯燥的中医术语变成可对话、可试玩的活体验，兼顾趣味与传播。"),
]
pos=[(0.9,2.5),(6.9,2.5),(0.9,4.0),(6.9,4.0),(0.9,5.5),(6.9,5.5)]
for (k,v),(l,t) in zip(items,pos):
    card(s,l,t,5.5,1.25,k,v)

# ---------------- 06 扫码试玩（末页 · 二维码居中） ----------------
s = prs.slides.add_slide(BLANK); bg(s, PAPER2); credit(s)
tag(s,"TRY / 扫码试玩")
heading(s,"扫码即玩《游医天下》")
qr = segno.make(PLAY_URL, error="h")
qr.save(QR_PATH, scale=12, border=2, dark="#2b2723", light="#ece7da")
QS = 4.0
ql = (13.333 - QS) / 2
s.shapes.add_picture(QR_PATH, Inches(ql), Inches(2.15), Inches(QS), Inches(QS))
set_text(box(s,1.0,6.45,11.333,0.5).text_frame,
         "手机扫码进入在线试玩 · 与「外泌体晨光」对话，体验望色归脏拟方", 15, INK2, False, PP_ALIGN.CENTER, 1)
set_text(box(s,1.0,6.95,11.333,0.4).text_frame, PLAY_URL, 11, GOLD, True, PP_ALIGN.CENTER, 1, "Arial")

out="/Users/xuyueqing/CodeBuddy/20260903210153/youyi-tianxia-3d/游医天下_游戏介绍.pptx"
prs.save(out)
print("saved", out, "slides:", len(prs.slides._sldIdLst))
